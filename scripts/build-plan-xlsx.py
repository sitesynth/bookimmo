#!/usr/bin/env python3
"""Regenerate public/plan/bookimmo_business_plan_v2.xlsx from public/plan/plan-data.json.

Single source of truth: plan-data.json. The deck (/plan) and text plan (/plan/text)
compute the same funnel at runtime via public/plan/model.js; this script mirrors that
logic as live Excel formulas so drivers can be changed inside the workbook.

Usage: python3 scripts/build-plan-xlsx.py
"""
import json
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / 'public' / 'plan' / 'plan-data.json').read_text())
OUT = ROOT / 'public' / 'plan' / 'bookimmo_business_plan_v2.xlsx'

MONTHS = DATA['months']
N = len(MONTHS)
COLS = [L(2 + i) for i in range(N)]          # B..M
YEAR_COL = L(2 + N)                           # N

hdr = Font(bold=True, size=11)
title_f = Font(bold=True, size=14)
sub_f = Font(italic=True, size=10, color='666666')
input_f = Font(color='FF1F4E79')
sec_fill = PatternFill('solid', fgColor='FFE8F0E8')
tot_fill = PatternFill('solid', fgColor='FFD9E8D9')
inp_fill = PatternFill('solid', fgColor='FFFFF7E0')
hdr_fill = PatternFill('solid', fgColor='FFDDE8F0')

wb = Workbook()

# ───────────────────────── Business Model (narrative) ─────────────────────────
ws = wb.active
ws.title = 'Business Model'
ws.column_dimensions['A'].width = 34
ws.column_dimensions['B'].width = 100
rows = [
    ('book.immo - Business Model v2', ''),
    ('GEO', DATA['meta']['geo'] + ' - стартовый рынок, затем другие города Германии'),
    ('', ''),
    ('ТЕЗИС ===', ''),
    ('Эксклюзивы до агрегаторов', 'CRM агентства выгружает на IS24 то, что не ушло по своей базе. Хорошие объекты на витрину не доходят. book.immo получает их из CRM до публикации - бесплатно для агентства, за сплит комиссии при сделке.'),
    ('Агентствам нужны хорошие лиды', 'Мы монетизируем не трафик, а квалификацию: агентство видит только проверенных кандидатов под конкретный эксклюзив.'),
    ('Каталог IS24 / Immowelt', 'Парсинг - для полноты каталога и SEO-трафика, не для сделок: там остатки.'),
    ('', ''),
    ('КАНАЛЫ ЛИДОВ ===', ''),
] + [(c['name'], f"[квалификация {int(c['qualifiedRate']*100)}%] {c['desc']}") for c in DATA['channels']] + [
    ('', ''),
    ('ВОРОНКА ===', ''),
] + [(f"{i+1}. {f['step']}", f['desc']) for i, f in enumerate(DATA['funnel'])] + [
    ('', ''),
    ('LEGAL ===', ''),
    ('Wohnungsvermittlungsgesetz', 'Нельзя брать с арендатора комиссию за посредничество. B2C fee - плата за софт (автозаявки, мониторинг, документы). Нужен письменный legal opinion до запуска.'),
    ('Bestellerprinzip', 'Комиссию платит собственник маклеру. Сплит book.immo - из B2B-агентского договора. Легально.'),
    ('GDPR', 'DPA с каждым агентством; consent клиента на передачу заявки.'),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(row=i, column=1, value=a)
    ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical='top')
    if a.endswith('==='):
        ws.cell(row=i, column=1).font = hdr
        ws.cell(row=i, column=1).fill = sec_fill
    elif i == 1:
        ws.cell(row=i, column=1).font = title_f

# ───────────────────────── Pre-launch budget ─────────────────────────
ws2 = wb.create_sheet('Budget Apr-Aug (USD)')
ws2.column_dimensions['A'].width = 46
bm = DATA['budget']['months']
ws2['A1'] = f"book.immo - Pre-Launch Budget ({bm[0]}–{bm[-1]} 2026, USD)"
ws2['A1'].font = title_f
m = DATA['meta']
ws2['A2'] = f"Инвестиция ${m['investmentUSD']:,}: ${m['spentUSD']:,} на запуск, ${m['reserveUSD']:,} резерв."
ws2['A2'].font = sub_f
ws2.cell(row=4, column=1, value='Category').font = hdr
for i, mo in enumerate(bm):
    ws2.cell(row=4, column=2 + i, value=mo).font = hdr
tot_c = len(bm) + 2
ws2.cell(row=4, column=tot_c, value='TOTAL').font = hdr
for c in range(1, tot_c + 1):
    ws2.cell(row=4, column=c).fill = hdr_fill
r = 5
for row in DATA['budget']['rows']:
    ws2.cell(row=r, column=1, value=row['name'])
    for i, v in enumerate(row['vals']):
        ws2.cell(row=r, column=2 + i, value=v)
    ws2.cell(row=r, column=tot_c, value=f'=SUM(B{r}:{L(len(bm)+1)}{r})')
    r += 1
ws2.cell(row=r, column=1, value='TOTAL')
for c in range(2, tot_c + 1):
    ws2.cell(row=r, column=c, value=f'=SUM({L(c)}5:{L(c)}{r-1})')
for c in range(1, tot_c + 1):
    ws2.cell(row=r, column=c).font = hdr
    ws2.cell(row=r, column=c).fill = tot_fill
r += 2
ws2.cell(row=r, column=1, value='Что построено:').font = hdr
for b in DATA['budget']['built']:
    r += 1
    ws2.cell(row=r, column=1, value='• ' + b)

# ───────────────────────── Scenario sheets (live formulas) ─────────────────────────
def scenario_sheet(key):
    sc = DATA['scenarios'][key]
    d = sc['drivers']
    ue = DATA['unitEconomics']
    ch = {c['key']: c for c in DATA['channels']}
    ws = wb.create_sheet(f"{sc['label']} - {DATA['meta']['geo']}")
    ws.column_dimensions['A'].width = 44
    for i in range(2, N + 3):
        ws.column_dimensions[L(i)].width = 10

    ws['A1'] = f"book.immo - {sc['label']}, {DATA['meta']['geo']}, запуск {DATA['meta']['launchMonth']} (EUR)"
    ws['A1'].font = title_f
    ws['A2'] = 'Жёлтые ячейки - входные драйверы (можно менять). Остальное считается формулами.'
    ws['A2'].font = sub_f
    ws.cell(row=4, column=1, value='')
    for i, mo in enumerate(MONTHS):
        ws.cell(row=4, column=2 + i, value=mo)
    ws.cell(row=4, column=2 + N, value='YEAR')
    for c in range(1, N + 3):
        ws.cell(row=4, column=c).font = hdr
        ws.cell(row=4, column=c).fill = hdr_fill

    row = [5]
    ref = {}

    def section(title):
        ws.cell(row=row[0], column=1, value=title).font = hdr
        ws.cell(row=row[0], column=1).fill = sec_fill
        row[0] += 1

    def put(name, key, values=None, formula=None, is_input=False, total=None, bold=False):
        r = row[0]
        ref[key] = r
        ws.cell(row=r, column=1, value=name)
        for i in range(N):
            c = ws.cell(row=r, column=2 + i)
            if values is not None:
                c.value = values[i] if isinstance(values, list) else values
            else:
                c.value = formula(COLS[i], i)
            if is_input:
                c.fill = inp_fill
                c.font = input_f
        if total == 'sum':
            ws.cell(row=r, column=2 + N, value=f'=SUM(B{r}:{COLS[-1]}{r})')
        elif total == 'avg':
            ws.cell(row=r, column=2 + N, value=f'=AVERAGE(B{r}:{COLS[-1]}{r})')
        if bold:
            for c in range(1, N + 3):
                ws.cell(row=r, column=c).font = hdr
                ws.cell(row=r, column=c).fill = tot_fill
        row[0] += 1
        return r

    def scalar(name, key, value):
        r = row[0]
        ref[key] = r
        ws.cell(row=r, column=1, value=name)
        c = ws.cell(row=r, column=2, value=value)
        c.fill = inp_fill
        c.font = input_f
        row[0] += 1

    def R(key, col):
        return f'{col}{ref[key]}'

    def S(key):
        return f'$B${ref[key]}'

    section('ДРАЙВЕРЫ - каналы (входные)')
    put('Paid: бюджет, €/мес', 'paid', d['paidSpend'], is_input=True, total='sum')
    scalar('Paid: CPC, €', 'cpc', d['cpc'])
    scalar('Paid: визит → регистрация', 'paidSignup', d['paidSignupRate'])
    put('Комьюнити: регистраций/мес', 'community', d['communitySignups'], is_input=True, total='sum')
    put('Партнёры (HR / вузы), шт.', 'partners', d['partners'], is_input=True)
    scalar('Лидов на партнёра / мес', 'leadsPerPartner', d['leadsPerPartner'])
    put('SEO: визитов/мес', 'seo', d['seoVisits'], is_input=True, total='sum')
    scalar('SEO: визит → регистрация', 'seoSignup', d['seoSignupRate'])
    put('Агентства: холодный BD (накопл.)', 'agencies', d['activeAgencies'], is_input=True)
    scalar('Визитов от агентства / мес', 'agVisits', d['agencyVisitsPerAgency'])
    scalar('Агентства: визит → регистрация', 'agSignup', d['agencySignupRate'])
    scalar('Referral: лидов на сделку прошлого месяца', 'refK', d['referralPerDeal'])
    scalar('Тёплый вход: доля self-deals -> партнёр-агентство', 'warmRate', d.get('catalogAgencyConversionRate', 0))
    row[0] += 1

    section('КВАЛИФИКАЦИЯ по каналу (доля регистраций, прошедших скрининг)')
    for k in ['paid', 'community', 'partners', 'seo', 'referral', 'agencies']:
        scalar(f"  {ch[k]['name']}", f'q_{k}', ch[k]['qualifiedRate'])
    row[0] += 1

    section('КОНВЕРСИЯ в сделку')
    scalar('Квалифицированный → сделка через агентство', 'toAg', d['qualToAgencyDeal'])
    scalar('Квалифицированный → сделка на каталоге IS24', 'toSelf', d['qualToSelfDeal'])
    put('Потолок сделок на агентство / мес', 'cap', d['dealsCapPerAgency'], is_input=True)
    row[0] += 1

    section('UNIT ECONOMICS')
    scalar('Сплит аренда, €/сделка', 'rent', ue['rentalSplitEUR'])
    scalar('Сплит продажа, €/сделка', 'sale', ue['saleSplitEUR'])
    scalar('Доля аренды', 'pctRent', ue['pctRental'])
    scalar('B2C fee, €', 'fee', ue['b2cFeeEUR'])
    scalar('Доля агентских сделок, платящих B2C fee', 'feeShare', ue['feeShareOnAgencyDeals'])
    scalar('Продажник: % от выручки агентских сделок', 'bdPct', ue.get('bdCommissionOnAgencyRevenue', 0))
    scalar('Zolak: средний чек мебели, €', 'furnTicket', ue.get('furnitureAvgTicketEUR', 0))
    scalar('Zolak: наша доля от мебели', 'furnPct', ue.get('furnitureCommissionPct', 0))
    put('Upsell агентствам, €/мес', 'upsell', d['upsells'], is_input=True, total='sum')
    row[0] += 1

    section('ВОРОНКА')
    # placeholder: cumulative warm-onboarded agencies (filled after selfDeals is defined below)
    warm_row = put('Тёплые агентства (накопл., от прошлых self-deals)', 'cumWarm', values=0)
    put('Эффективные агентства (холодные + тёплые)', 'effAg', formula=lambda c, i: f'={R("agencies", c)}+{R("cumWarm", c)}', bold=False)
    put('Регистрации: paid', 's_paid', formula=lambda c, i: f'={R("paid", c)}/{S("cpc")}*{S("paidSignup")}', total='sum')
    put('Регистрации: комьюнити', 's_comm', formula=lambda c, i: f'={R("community", c)}', total='sum')
    put('Регистрации: SEO', 's_seo', formula=lambda c, i: f'={R("seo", c)}*{S("seoSignup")}', total='sum')
    put('Регистрации: агентства', 's_ag', formula=lambda c, i: f'={R("effAg", c)}*{S("agVisits")}*{S("agSignup")}', total='sum')
    put('Лиды: партнёры (уже квалифицированы)', 's_part', formula=lambda c, i: f'={R("partners", c)}*{S("leadsPerPartner")}', total='sum')
    put('Итого регистраций', 'signups', formula=lambda c, i: f'={R("s_paid", c)}+{R("s_comm", c)}+{R("s_seo", c)}+{R("s_ag", c)}+{R("s_part", c)}', total='sum', bold=True)
    put('Квалифицировано: paid', 'q1', formula=lambda c, i: f'={R("s_paid", c)}*{S("q_paid")}', total='sum')
    put('Квалифицировано: комьюнити', 'q2', formula=lambda c, i: f'={R("s_comm", c)}*{S("q_community")}', total='sum')
    put('Квалифицировано: SEO', 'q3', formula=lambda c, i: f'={R("s_seo", c)}*{S("q_seo")}', total='sum')
    put('Квалифицировано: агентства', 'q4', formula=lambda c, i: f'={R("s_ag", c)}*{S("q_agencies")}', total='sum')
    put('Квалифицировано: партнёры', 'q5', formula=lambda c, i: f'={R("s_part", c)}*{S("q_partners")}', total='sum')
    # referral depends on previous month's deals → defined after deals rows; placeholder row now
    ref_row = put('Квалифицировано: referral', 'q6', values=0, total='sum')
    put('Итого квалифицированных', 'qual', formula=lambda c, i: f'={R("q1", c)}+{R("q2", c)}+{R("q3", c)}+{R("q4", c)}+{R("q5", c)}+{R("q6", c)}', total='sum', bold=True)
    put('Сделки через агентства', 'agDeals', formula=lambda c, i: f'=ROUND(MIN({R("qual", c)}*{S("toAg")},{R("effAg", c)}*{R("cap", c)}),0)', total='sum')
    put('Сделки на каталоге IS24', 'selfDeals', formula=lambda c, i: f'=ROUND({R("qual", c)}*{S("toSelf")},0)', total='sum')
    put('Сделок всего', 'deals', formula=lambda c, i: f'={R("agDeals", c)}+{R("selfDeals", c)}', total='sum', bold=True)
    # fill referral formulas now that deals row exists
    for i, c in enumerate(COLS):
        cell = ws.cell(row=ref_row, column=2 + i)
        cell.value = 0 if i == 0 else f'={COLS[i-1]}{ref["deals"]}*{S("refK")}*{S("q_referral")}'
    # fill warm-agency cumulative now that selfDeals is defined: this month's warm = prev cumWarm + prev selfDeals*rate
    for i, c in enumerate(COLS):
        cell = ws.cell(row=warm_row, column=2 + i)
        cell.value = 0 if i == 0 else f'={COLS[i-1]}{ref["cumWarm"]}+ROUND({COLS[i-1]}{ref["selfDeals"]}*{S("warmRate")},0)'
    put('Сделок с B2C fee', 'feeDeals', formula=lambda c, i: f'={R("selfDeals", c)}+ROUND({R("agDeals", c)}*{S("feeShare")},0)', total='sum')
    row[0] += 1

    section('ВЫРУЧКА')
    put('R1+R2: сплит комиссии (аренда/продажа)', 'revAg', formula=lambda c, i: f'=ROUND({R("agDeals", c)}*({S("pctRent")}*{S("rent")}+(1-{S("pctRent")})*{S("sale")}),0)', total='sum')
    put('R3: B2C fee', 'revFee', formula=lambda c, i: f'={R("feeDeals", c)}*{S("fee")}', total='sum')
    put('R4: upsell', 'revUp', formula=lambda c, i: f'={R("upsell", c)}', total='sum')
    put('R5: реферал на мебель (Zolak)', 'revFurn', formula=lambda c, i: f'=ROUND({R("deals", c)}*{S("furnTicket")}*{S("furnPct")},0)', total='sum')
    put('TOTAL REVENUE', 'rev', formula=lambda c, i: f'={R("revAg", c)}+{R("revFee", c)}+{R("revUp", c)}+{R("revFurn", c)}', total='sum', bold=True)
    row[0] += 1

    section('РАСХОДЫ')
    put('Маркетинг: paid', 'cPaid', formula=lambda c, i: f'={R("paid", c)}', total='sum')
    cost_keys = []
    for j, cr in enumerate(sc['costs']):
        k = f'c{j}'
        cost_keys.append(k)
        put(cr['name'], k, cr['vals'], is_input=True, total='sum')
    put('Продажник: комиссия (% от R1+R2)', 'bdComm', formula=lambda c, i: f'=ROUND({R("revAg", c)}*{S("bdPct")},0)', total='sum')
    cost_keys.append('bdComm')
    put('TOTAL COSTS', 'cost', formula=lambda c, i: '=' + '+'.join([R('cPaid', c)] + [R(k, c) for k in cost_keys]), total='sum', bold=True)
    row[0] += 1

    put('EBITDA', 'ebitda', formula=lambda c, i: f'={R("rev", c)}-{R("cost", c)}', total='sum', bold=True)
    put(f"Кэш (старт €{DATA['meta']['reserveEUR']:,})", 'cash', formula=lambda c, i: (f"={DATA['meta']['reserveEUR']}+{R('ebitda', c)}" if i == 0 else f'={COLS[i-1]}{ref["cash"]}+{R("ebitda", c)}'))
    row[0] += 1

    section('ЮНИТ-ЭКОНОМИКА (год)')
    yr = YEAR_COL
    ws.cell(row=row[0], column=1, value='CAC на квалифицированного лида, €')
    ws.cell(row=row[0], column=2, value=f'=ROUND({yr}{ref["paid"]}/{yr}{ref["qual"]},0)'); row[0] += 1
    ws.cell(row=row[0], column=1, value='CAC на сделку, €')
    ws.cell(row=row[0], column=2, value=f'=ROUND({yr}{ref["paid"]}/{yr}{ref["deals"]},0)'); row[0] += 1
    ws.cell(row=row[0], column=1, value='Выручка на агентскую сделку, €')
    ws.cell(row=row[0], column=2, value=f'=ROUND({S("pctRent")}*{S("rent")}+(1-{S("pctRent")})*{S("sale")},0)'); row[0] += 1
    ws.cell(row=row[0], column=1, value='LTV / CAC (агентский канал)')
    ws.cell(row=row[0], column=2, value=f'=ROUND(B{row[0]-1}/B{row[0]-2},1)'); row[0] += 1

for key in ['base', 'pessimistic']:
    scenario_sheet(key)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)


def fix_app_xml(path, sheet_names):
    """openpyxl omits docProps/app.xml's HeadingPairs/TitlesOfParts (the sheet-name
    manifest some Excel builds cross-check against xl/workbook.xml's <sheets>).
    Rewrite that one part in place with proper values."""
    ns_vt = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
    titles = ''.join(f'<vt:lpstr>{n}</vt:lpstr>' for n in sheet_names)
    app_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        f'xmlns:vt="{ns_vt}">'
        '<Application>Microsoft Excel</Application>'
        '<DocSecurity>0</DocSecurity>'
        '<ScaleCrop>false</ScaleCrop>'
        '<HeadingPairs>'
        f'<vt:vector size="2" baseType="variant">'
        f'<vt:variant><vt:lpstr>Worksheets</vt:lpstr></vt:variant>'
        f'<vt:variant><vt:i4>{len(sheet_names)}</vt:i4></vt:variant>'
        f'</vt:vector>'
        '</HeadingPairs>'
        f'<TitlesOfParts><vt:vector size="{len(sheet_names)}" baseType="lpstr">{titles}</vt:vector></TitlesOfParts>'
        '<LinksUpToDate>false</LinksUpToDate>'
        '<SharedDoc>false</SharedDoc>'
        '<HyperlinksChanged>false</HyperlinksChanged>'
        '<AppVersion>16.0300</AppVersion>'
        '</Properties>'
    )
    tmp = path.with_suffix('.tmp.xlsx')
    with zipfile.ZipFile(path, 'r') as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = app_xml.encode('utf-8') if item.filename == 'docProps/app.xml' else zin.read(item.filename)
            zout.writestr(item, data)
    shutil.move(tmp, path)


fix_app_xml(OUT, [ws.title for ws in wb.worksheets])
print(f'Saved {OUT.relative_to(ROOT)}')
